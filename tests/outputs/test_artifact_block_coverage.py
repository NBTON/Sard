"""Canonical block coverage: timeline/sources/page-break/list render in all formats.

Regression: renderers silently dropped non-core block types (timeline,
sources, page-break coerced to "item" or ignored). Every canonical type in
document.BLOCK_TYPES must survive HTML/PDF/DOCX/PPTX with its text intact.
"""
from __future__ import annotations

import io

from pypdf import PdfReader


def _canonical_doc():
    from sard.outputs.document import ArtifactDocument
    from sard.outputs.orchestrator import ArtifactRequest

    req = ArtifactRequest(
        format="pdf",
        kind="document",
        title="تقرير الكتل الأساسية",
        topic="الدرعية",
        raw_text="نص توثيقي.",
        content_data={
            "sections": [
                {
                    "id": "s1",
                    "title": "الجدول الزمني",
                    "blocks": [
                        {
                            "id": "b1",
                            "type": "timeline",
                            "text": "أحداث الدرعية",
                            "source_ids": ["CIT-001"],
                            "data": {
                                "items": [
                                    {"date": "1727", "title": "تأسيس الدرعية", "text": "عاصمة الدولة الأولى"},
                                    {"date": "1765", "title": "قصر سلوى", "text": "مقر الحكم"},
                                ]
                            },
                        },
                        {"id": "b2", "type": "list", "text": "بند موثق أول", "source_ids": ["CIT-001"]},
                        {"id": "b3", "type": "list", "text": "بند موثق ثان", "source_ids": ["CIT-001"]},
                        {"id": "b4", "type": "page-break", "text": ""},
                    ],
                },
                {
                    "id": "s2",
                    "title": "المصادر",
                    "blocks": [
                        {
                            "id": "b5",
                            "type": "sources",
                            "text": "دارة الملك عبد العزيز",
                            "source_ids": ["CIT-001"],
                            "data": {"items": [{"title": "دارة الملك عبد العزيز", "url": "https://example.com"}]},
                        },
                    ],
                },
            ]
        },
        sources=({"citation_id": "CIT-001", "title": "دارة الملك عبد العزيز", "url": "https://example.com"},),
        metadata={"artifact_id": "art-blockcov", "run_id": "run-blockcov", "version": 1},
    )
    doc = ArtifactDocument.from_request(req)
    types = [b.block_type for s in doc.sections for b in s.blocks]
    assert "timeline" in types and "sources" in types and "page-break" in types and "list" in types
    return doc


def test_canonical_blocks_render_html():
    from sard.outputs.html import render_html_document

    out = render_html_document(_canonical_doc())
    assert "1727" in out and "تأسيس الدرعية" in out  # timeline kept
    assert "sard-timeline" in out
    assert "دارة الملك عبد العزيز" in out  # sources kept
    assert "sard-page-break" in out
    assert "بند موثق أول" in out  # list kept


def test_canonical_blocks_render_pdf():
    from sard.outputs.pdf import build_pdf_from_document

    data = build_pdf_from_document(_canonical_doc())
    assert data.startswith(b"%PDF")
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) >= 2  # page-break honored
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    assert "1727" in text  # timeline kept


def test_canonical_blocks_render_docx():
    from docx import Document

    from sard.outputs.office_docx import DocxGenerator

    data = DocxGenerator().build_from_document(_canonical_doc())
    assert data.startswith(b"PK")
    paras = Document(io.BytesIO(data)).paragraphs
    full = "\n".join(p.text for p in paras)
    assert "1727" in full and "تأسيس الدرعية" in full
    assert "دارة الملك عبد العزيز" in full


def test_canonical_blocks_render_pptx():
    from pptx import Presentation

    from sard.outputs.office import PresentationGenerator

    data = PresentationGenerator().build_from_document(_canonical_doc())
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) >= 3  # title + content across the page-break
    full = "\n".join(
        shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "1727" in full
